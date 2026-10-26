"""Сборка сдаваемых документов СРС (docx) из Markdown-исходников в DV/src.

Запуск: pip install python-docx && python DV/build_docs.py
Шапка каждого документа — по образцу разбора СРС из материалов практики 2.5.
Поддерживается подмножество Markdown: заголовки ##/###, абзацы, списки, таблицы,
**жирный**, *курсив*, `код`, изображения ![подпись](путь).
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

DV = Path(__file__).resolve().parent
COURSE = "Интеллектуальный анализ и визуализация данных"
TRACK = "ML-инжиниринг"
TEAM = "Храмов Владислав, Гарбуз Александр, Семенов Илья"
GROUP = "____________"  # впишите номер группы

DOCS = {
    "srs_1.3_dikw.md": "IAVD_SRS_1.3_DIKW.docx",
    "srs_2.1_encodings.md": "IAVD_SRS_2.1_kodirovki_eskiz.docx",
    "srs_2.2_design.md": "IAVD_SRS_2.2_Enskomb_do_posle.docx",
    "eda_bank.md": "IAVD_1.5_EDA_bank_otvety.docx",
    "srs_3.1_metrics.md": "IAVD_SRS_3.1_metriki.docx",
}

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+`)")


def add_runs(par, text: str):
    for part in INLINE.split(text.replace("\\|", "|")):
        if not part:
            continue
        if part.startswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Courier New"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def split_row(line: str) -> list[str]:
    cells = re.split(r"(?<!\\)\|", line.strip().strip("|"))
    return [c.strip() for c in cells]


def front_matter(text: str) -> tuple[dict, list[str]]:
    lines = text.splitlines()
    meta = {}
    if lines and lines[0] == "---":
        end = lines.index("---", 1)
        for ln in lines[1:end]:
            key, _, value = ln.partition(":")
            meta[key.strip()] = value.strip()
        lines = lines[end + 1:]
    return meta, lines


def header(doc: Document, meta: dict):
    p = doc.add_paragraph()
    run = p.add_run(meta["title"])
    run.bold = True
    run.font.size = Pt(16)
    if meta.get("subtitle"):
        doc.add_paragraph().add_run(meta["subtitle"]).italic = True
    for label, value in [
        ("Курс", COURSE),
        ("Модуль", meta.get("module", "")),
        ("Трек", TRACK),
        ("Выполнили", TEAM),
        ("Группа", GROUP),
        ("Датасет", meta.get("dataset", "")),
    ]:
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(0)
        par.add_run(f"{label}: ").bold = True
        par.add_run(value)
    doc.add_paragraph()


def build(src: Path, out: Path):
    meta, lines = front_matter(src.read_text(encoding="utf-8"))
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2)
    header(doc, meta)

    i = 0
    while i < len(lines):
        ln = lines[i]
        img = re.match(r"^!\[(.*)\]\((.+)\)$", ln.strip())
        if img:
            caption, path = img.groups()
            doc.add_picture(str((src.parent / path).resolve()), width=Cm(16.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = cap.add_run(caption)
            run.italic = True
            run.font.size = Pt(10)
        elif ln.startswith("#"):
            level = len(ln) - len(ln.lstrip("#"))
            doc.add_heading(ln.lstrip("#").strip(), level=min(level - 1, 3) or 1)
        elif ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[\s\-|:]+\|$", lines[i]):
                    rows.append(split_row(lines[i]))
                i += 1
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    par = table.cell(r, c).paragraphs[0]
                    add_runs(par, value)
                    for run in par.runs:
                        run.font.size = Pt(10.5)
                        if r == 0:
                            run.bold = True
            doc.add_paragraph()
            continue
        elif re.match(r"^(- |\d+\. )", ln):
            style = "List Bullet" if ln.startswith("- ") else "List Number"
            add_runs(doc.add_paragraph(style=style), re.sub(r"^(- |\d+\. )", "", ln))
        elif ln.strip():
            add_runs(doc.add_paragraph(), ln)
        i += 1
    doc.save(out)


if __name__ == "__main__":
    for name, out in DOCS.items():
        build(DV / "src" / name, DV / out)
        print(f"{name} → {out}")
